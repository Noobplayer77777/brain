from pptx import Presentation

def parse_pptx(file_path: str) -> list[dict]:
    prs = Presentation(file_path)
    slides_data = []
    
    for i, slide in enumerate(prs.slides):
        text_parts = []
        
        # Extract from shapes
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text_parts.append(shape.text.strip())
                
        # Extract from notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                text_parts.append("Speaker Notes: " + notes)
                
        slides_data.append({
            "slide_no": i + 1,
            "text": "\n".join(text_parts).strip()
        })
        
    return slides_data
