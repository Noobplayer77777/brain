import os
import fitz
from pptx import Presentation
from PIL import Image, ImageDraw

os.makedirs("data/raw/DBMS", exist_ok=True)

# 1. Real text PDF
doc = fitz.open()
page = doc.new_page()
page.insert_text((50, 50), "DBMS First Normal Form (1NF).")
page.insert_text((50, 70), "A relation is in 1NF if it contains only atomic values.")
doc.save("data/raw/DBMS/dbms_notes.pdf")

# 2. Real PPTX
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "ACID Properties"
slide.placeholders[1].text = "Atomicity, Consistency, Isolation, Durability."
prs.save("data/raw/DBMS/dbms_acid.pptx")

# 3. Scanned PDF
img = Image.new('RGB', (400, 200), color=(255,255,255))
d = ImageDraw.Draw(img)
d.text((10,10), "Scanned DBMS query: SELECT * FROM users;", fill=(0,0,0))
img.save("data/raw/DBMS/scanned_query.pdf", "PDF")

print("Test data created.")
